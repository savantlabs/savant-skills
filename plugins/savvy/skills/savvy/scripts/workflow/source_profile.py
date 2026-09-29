#!/usr/bin/env python3
"""Profile local sources for Savant Planner/Builder handoffs.

The profiler is intentionally local and API-free. It turns user-supplied CSV/Excel/PDF files into
bounded, reusable evidence. Tabular files get observed schema, sample values, date-standardization
hints, likely keys, and cross-source key overlap. PDFs get a binary-document source suggestion so
they do not fall through to CSV parsing. Agents should use this before hand-authoring source plans.
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import json
import re
import sys
from collections import defaultdict
from pathlib import Path
from typing import Any

SCRIPT_DIR = Path(__file__).resolve().parents[1]
if str(SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPT_DIR))

from savant_api.fileio import save_json, workspace_tmp  # noqa: E402


DATE_NAME_RE = re.compile(r"(date|dt|timestamp|time)", re.IGNORECASE)
KEY_NAME_RE = re.compile(r"\b(id|key|code|number|num|no|#|ref|reference|document|doc|vendor|account|item|rct|rtv|profit ctr)\b", re.IGNORECASE)
KEY_SUPPRESS_RE = re.compile(
    r"(date|timestamp|time|amount|total|count|status|meaning|matched|variance|flag|description|text|note|period)",
    re.IGNORECASE,
)
IDENTIFIER_NAME_RE = re.compile(r"(id|#|num|number|row|reference|ref|document|doc)", re.IGNORECASE)
HEADER_LABEL_RE = re.compile(
    r"\b(id|type|date|status|source|amount|total|count|vendor|account|reference|ref|document|description|"
    r"population|support|matched|variance|note|class|period|profit|disposition|meaning|groups)\b",
    re.IGNORECASE,
)
DEFAULT_SCAN_ROWS = 50_000
DEFAULT_SAMPLE_VALUES = 5
OPENPYXL_DEPENDENCY = "openpyxl>=3.1,<4"


def _load_openpyxl():
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError(f"Excel profiling requires Python package `{OPENPYXL_DEPENDENCY}`.") from exc
    return openpyxl


def _json_safe(value: Any) -> Any:
    if isinstance(value, (dt.datetime, dt.date)):
        return value.isoformat()
    return value


def _blank(value: Any) -> bool:
    return value is None or str(value).strip() == ""


def _trim(value: Any) -> str:
    return str(value).strip()


def _try_float(value: Any) -> float | None:
    if _blank(value):
        return None
    try:
        return float(str(value).strip().replace(",", ""))
    except (TypeError, ValueError):
        return None


def _try_int(value: Any) -> int | None:
    number = _try_float(value)
    if number is None or not float(number).is_integer():
        return None
    return int(number)


def infer_type(values: list[Any]) -> str:
    nonempty = [v for v in values if not _blank(v)]
    if not nonempty:
        return "string"
    if all(isinstance(v, bool) for v in nonempty):
        return "boolean"
    if all(isinstance(v, int) and not isinstance(v, bool) for v in nonempty):
        return "integer"
    if all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in nonempty):
        return "number"
    if all(_try_int(v) is not None for v in nonempty):
        return "integer"
    if all(_try_float(v) is not None for v in nonempty):
        return "number"
    if all(isinstance(v, dt.datetime) for v in nonempty):
        return "datetime"
    if all(isinstance(v, dt.date) for v in nonempty):
        return "date"
    return "string"


def _parse_date_string(value: Any) -> dt.date | None:
    if _blank(value):
        return None
    text = str(value).strip()
    if not re.search(r"\d", text):
        return None
    for fmt in (
        "%Y-%m-%d",
        "%Y/%m/%d",
        "%m/%d/%Y",
        "%m-%d-%Y",
        "%d/%m/%Y",
        "%d-%m-%Y",
        "%Y-%m-%d %H:%M:%S",
        "%m/%d/%Y %H:%M:%S",
    ):
        try:
            parsed = dt.datetime.strptime(text, fmt)
            return parsed.date()
        except ValueError:
            pass
    try:
        return dt.datetime.fromisoformat(text.replace("Z", "+00:00")).date()
    except ValueError:
        return None


def _excel_serial_date_stats(values: list[Any]) -> tuple[bool, bool, int, int]:
    nonempty = [v for v in values if not _blank(v)]
    if not nonempty:
        return False, False, 0, 0
    numbers = [_try_float(v) for v in nonempty]
    # Covers roughly 1954-2064 in Excel serial days; broad enough for business data.
    serials = [number for number in numbers if number is not None and 20_000 <= float(number) <= 60_000]
    if not serials:
        return False, False, 0, len(nonempty)
    ratio = len(serials) / len(nonempty)
    has_time_fraction = any(not float(number).is_integer() for number in serials)
    return ratio >= 0.6, has_time_fraction, len(serials), len(nonempty)


def date_candidate(name: str, values: list[Any], data_type: str) -> dict[str, Any] | None:
    nonempty = [v for v in values if not _blank(v)]
    if not nonempty:
        return None
    name_suggests_date = bool(DATE_NAME_RE.search(name))
    if data_type in {"date", "datetime"}:
        return {
            "column": name,
            "observedDataType": data_type,
            "expectedDataType": data_type,
            "needsAdapter": False,
            "reason": "source values are already date/datetime typed",
        }
    is_excel_serial, has_time_fraction, serial_count, nonempty_count = _excel_serial_date_stats(nonempty)
    if name_suggests_date and is_excel_serial:
        return {
            "column": name,
            "observedDataType": data_type,
            "expectedDataType": "datetime" if has_time_fraction else "date",
            "needsAdapter": True,
            "reason": "date-like column name with Excel serial-number values"
            + (" including time fractions" if has_time_fraction else "")
            + (f" ({serial_count}/{nonempty_count} non-empty values)" if serial_count != nonempty_count else ""),
        }
    parsed = [_parse_date_string(v) for v in nonempty[:100]]
    parsed_count = sum(1 for value in parsed if value is not None)
    if name_suggests_date and parsed and parsed_count / len(parsed) >= 0.6:
        return {
            "column": name,
            "observedDataType": data_type,
            "expectedDataType": "date",
            "needsAdapter": True,
            "reason": f"date-like column name with parseable date strings ({parsed_count}/{len(parsed)} sampled values)",
        }
    return None


def _sample_values(values: list[Any], limit: int) -> list[Any]:
    samples: list[Any] = []
    seen: set[str] = set()
    for value in values:
        if _blank(value):
            continue
        safe = _json_safe(value)
        key = json.dumps(safe, sort_keys=True, default=str)
        if key in seen:
            continue
        seen.add(key)
        samples.append(safe)
        if len(samples) >= limit:
            break
    return samples


def _excel_stored_value(cell, epoch):
    value = cell.value
    if value is None:
        return None
    if getattr(cell, "is_date", False) and isinstance(value, (dt.datetime, dt.date)):
        from openpyxl.utils.datetime import to_excel

        serial = to_excel(value, epoch)
        return int(serial) if float(serial).is_integer() else serial
    return value


def _normalize_row(row: list[Any], width: int) -> list[Any]:
    normalized = list(row[:width])
    if len(normalized) < width:
        normalized.extend([None] * (width - len(normalized)))
    return normalized


def _make_unique_columns(values: list[Any]) -> list[str]:
    columns: list[str] = []
    counts: defaultdict[str, int] = defaultdict(int)
    for index, value in enumerate(values):
        base = str(value).strip() if value is not None and str(value).strip() else f"col_{index + 1}"
        counts[base] += 1
        columns.append(base if counts[base] == 1 else f"{base}_{counts[base]}")
    return columns


CSV_FALLBACK_ENCODING = "cp1252"


def _open_text(path: Path) -> tuple[Any, str]:
    """Open a text file as UTF-8, falling back to cp1252 (a superset of ISO-8859-1) when the
    bytes are not valid UTF-8. Returns (handle, encoding_used)."""
    try:
        with path.open(encoding="utf-8-sig") as probe:
            probe.read()
        return path.open(encoding="utf-8-sig", newline=""), "utf-8"
    except UnicodeDecodeError:
        return path.open(encoding=CSV_FALLBACK_ENCODING, newline=""), CSV_FALLBACK_ENCODING


def _read_csv(path: Path, *, delimiter: str, max_rows: int) -> tuple[list[str], list[list[Any]], bool]:
    columns, rows, truncated, _encoding = _read_csv_with_encoding(path, delimiter=delimiter, max_rows=max_rows)
    return columns, rows, truncated


def _read_csv_with_encoding(path: Path, *, delimiter: str, max_rows: int) -> tuple[list[str], list[list[Any]], bool, str]:
    handle, encoding = _open_text(path)
    with handle:
        reader = csv.reader(handle, delimiter=delimiter)
        try:
            header = next(reader)
        except StopIteration as exc:
            raise ValueError(f"file has no rows: {path}") from exc
        columns = _make_unique_columns(header)
        rows: list[list[Any]] = []
        truncated = False
        for index, row in enumerate(reader):
            if index >= max_rows:
                truncated = True
                break
            rows.append(_normalize_row(row, len(columns)))
    return columns, rows, truncated, encoding


def _header_score(row: list[Any], following_rows: list[list[Any]]) -> float:
    nonempty_values = [value for value in row if not _blank(value)]
    if not nonempty_values:
        return 0.0
    string_count = sum(1 for value in nonempty_values if _try_float(value) is None)
    unique_count = len({_trim(value).lower() for value in nonempty_values})
    following_density = 0.0
    if following_rows:
        densities = []
        for next_row in following_rows:
            width = max(len(row), len(next_row), 1)
            densities.append(sum(1 for value in next_row if not _blank(value)) / width)
        following_density = sum(densities) / len(densities)
    return (
        len(nonempty_values) * 0.35
        + (string_count / len(nonempty_values)) * 2.0
        + (unique_count / len(nonempty_values)) * 1.5
        + following_density * 1.5
    )


def _detect_header_row(rows: list[list[Any]], *, search_limit: int = 25) -> int:
    if not rows:
        return 0
    best_index = 0
    best_score = -1.0
    for index, row in enumerate(rows[:search_limit]):
        score = _header_score(row, rows[index + 1:index + 4])
        if score > best_score:
            best_score = score
            best_index = index
    return best_index


def _read_xlsx(
    path: Path,
    sheet: str,
    *,
    max_rows: int,
    header_row: int | None = None,
    data_start_row: int | None = None,
    data_end_row: int | None = None,
) -> tuple[list[str], list[list[Any]], bool, int]:
    openpyxl = _load_openpyxl()
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if sheet not in workbook.sheetnames:
        raise ValueError(f"sheet {sheet!r} not found in {path}; available sheets: {workbook.sheetnames}")
    worksheet = workbook[sheet]
    raw_rows: list[list[Any]] = []
    truncated = False
    read_limit = max_rows + 50
    if data_end_row:
        read_limit = max(read_limit, data_end_row)
    for index, row_cells in enumerate(worksheet.iter_rows()):
        if index >= read_limit:
            truncated = True
            break
        raw_rows.append([_excel_stored_value(cell, workbook.epoch) for cell in row_cells])
    if not raw_rows:
        raise ValueError(f"sheet has no rows: {path}#{sheet}")
    if header_row is not None:
        if header_row <= 0:
            raise ValueError(f"header_row must be 1-based and positive for {path}#{sheet}.")
        header_index = header_row - 1
        if header_index >= len(raw_rows):
            raise ValueError(f"header_row {header_row} is beyond scanned rows for {path}#{sheet}.")
    else:
        header_index = _detect_header_row(raw_rows)
    columns = _make_unique_columns(raw_rows[header_index])
    start_index = (data_start_row - 1) if data_start_row else header_index + 1
    if start_index <= header_index:
        raise ValueError(f"data_start_row must be after header_row for {path}#{sheet}.")
    end_index = data_end_row if data_end_row else start_index + max_rows
    rows = [_normalize_row(row, len(columns)) for row in raw_rows[start_index:min(end_index, start_index + max_rows)]]
    if len(raw_rows) > min(end_index, start_index + max_rows):
        truncated = True
    return columns, rows, truncated, header_index + 1


def excel_sheet_names(path: Path) -> list[str]:
    openpyxl = _load_openpyxl()
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    return list(workbook.sheetnames)


def _load_hints(path: Path | None) -> list[dict[str, Any]]:
    if path is None:
        return []
    payload = json.loads(path.read_text(encoding="utf-8"))
    if isinstance(payload, dict):
        values = payload.get("sources") or payload.get("files") or []
    else:
        values = payload
    if not isinstance(values, list):
        raise ValueError("source profile hints must be a list or an object with `sources`.")
    hints = []
    for index, value in enumerate(values):
        if not isinstance(value, dict):
            raise ValueError(f"hints[{index}] must be an object.")
        hints.append(value)
    return hints


def _path_key(path: Path | str) -> str:
    return str(Path(path).expanduser().resolve())


def _hints_for_file(hints: list[dict[str, Any]], path: Path) -> list[dict[str, Any]]:
    key = _path_key(path)
    matches = []
    for hint in hints:
        hint_path = hint.get("path") or hint.get("file")
        if hint_path and _path_key(str(hint_path)) == key:
            source_hints = hint.get("sources")
            if isinstance(source_hints, list):
                for source_hint in source_hints:
                    if isinstance(source_hint, dict):
                        matches.append({**source_hint, "path": str(path)})
            else:
                matches.append(hint)
    return matches


def _hint_int(hint: dict[str, Any], *keys: str) -> int | None:
    for key in keys:
        value = hint.get(key)
        if value in (None, ""):
            continue
        try:
            parsed = int(value)
        except (TypeError, ValueError) as exc:
            raise ValueError(f"hint `{key}` must be an integer.") from exc
        if parsed <= 0:
            raise ValueError(f"hint `{key}` must be positive.")
        return parsed
    return None


def _column_values(rows: list[list[Any]], index: int) -> list[Any]:
    return [row[index] if index < len(row) else None for row in rows]


def _embedded_header_rows(rows: list[list[Any]], *, header_row: int, limit: int = 12) -> list[int]:
    embedded = []
    for offset, row in enumerate(rows):
        following = rows[offset + 1:offset + 4]
        if not following:
            continue
        nonempty = [value for value in row if not _blank(value)]
        if len(nonempty) < 2:
            continue
        string_count = sum(1 for value in nonempty if _try_float(value) is None)
        header_label_count = sum(1 for value in nonempty if _try_float(value) is None and HEADER_LABEL_RE.search(str(value)))
        numeric_count = len(nonempty) - string_count
        if string_count < 2 or header_label_count < 2 or numeric_count > 1:
            continue
        if _header_score(row, following) >= 3.0:
            embedded.append(header_row + 1 + offset)
            if len(embedded) >= limit:
                break
    return embedded


def column_profiles(columns: list[str], rows: list[list[Any]], *, sample_limit: int) -> tuple[list[dict], list[dict]]:
    observed: list[dict[str, Any]] = []
    date_candidates: list[dict[str, Any]] = []
    for index, name in enumerate(columns):
        values = _column_values(rows, index)
        data_type = infer_type(values)
        nonempty = [v for v in values if not _blank(v)]
        distinct = {_trim(v) for v in nonempty}
        profile = {
            "name": name,
            "dataType": data_type,
            "sampleValues": _sample_values(values, sample_limit),
            "nonEmptyCount": len(nonempty),
            "distinctCount": len(distinct),
        }
        observed.append(profile)
        candidate = date_candidate(name, values, data_type)
        if candidate:
            date_candidates.append(candidate)
    return observed, date_candidates


def expected_schema_from_profile(observed_schema: list[dict], date_candidates: list[dict]) -> list[dict]:
    date_types = {candidate["column"]: candidate["expectedDataType"] for candidate in date_candidates}
    return [
        {
            "name": column["name"],
            "dataType": date_types.get(column["name"], column["dataType"]),
            "mappedFrom": column["name"],
            "required": False,
        }
        for column in observed_schema
    ]


def _key_score(column: dict, row_count: int) -> float:
    if row_count <= 0:
        return 0.0
    name = str(column.get("name") or "")
    has_positive_name = bool(KEY_NAME_RE.search(name))
    if KEY_SUPPRESS_RE.search(name):
        return 0.0
    if not has_positive_name:
        return 0.0
    nonempty_ratio = column.get("nonEmptyCount", 0) / row_count
    distinct_ratio = column.get("distinctCount", 0) / max(column.get("nonEmptyCount", 0), 1)
    name_bonus = 0.35
    return min(1.0, distinct_ratio * 0.65 + nonempty_ratio * 0.25 + name_bonus)


def key_candidates(observed_schema: list[dict], row_count: int, *, limit: int = 12) -> list[dict]:
    candidates = []
    for column in observed_schema:
        score = _key_score(column, row_count)
        if score >= 0.55:
            candidates.append(
                {
                    "column": column["name"],
                    "score": round(score, 3),
                    "nonEmptyCount": column.get("nonEmptyCount", 0),
                    "distinctCount": column.get("distinctCount", 0),
                    "reason": "name/distinctness suggests an identifier or join key",
                }
            )
    return sorted(candidates, key=lambda item: (-item["score"], item["column"]))[:limit]


def _source_values(source: dict, column: str) -> list[Any]:
    index = source["_columns"].index(column)
    return _column_values(source["_rows"], index)


def _column_profile(source: dict, column: str) -> dict[str, Any]:
    for profile in source.get("observed_schema", []):
        if isinstance(profile, dict) and profile.get("name") == column:
            return profile
    return {}


def _unique_ratio(source: dict, column: str) -> float:
    profile = _column_profile(source, column)
    nonempty = profile.get("nonEmptyCount") or 0
    distinct = profile.get("distinctCount") or 0
    return (distinct / nonempty) if nonempty else 0.0


def _norm_exact(value: Any) -> str:
    return _trim(value).upper()


def _norm_alnum(value: Any) -> str:
    return re.sub(r"[^A-Z0-9]", "", _norm_exact(value))


def _norm_suffix_digits(value: Any) -> str:
    match = re.search(r"(\d{4,})$", _trim(value))
    return match.group(1) if match else ""


def _mostly_digits(values: list[Any]) -> bool:
    nonempty = [value for value in values if not _blank(value)]
    if not nonempty:
        return False
    digit_values = [value for value in nonempty if re.fullmatch(r"\d+", _trim(value))]
    return len(digit_values) / len(nonempty) >= 0.8


def _has_prefixed_trailing_digits(values: list[Any]) -> bool:
    nonempty = [value for value in values if not _blank(value)]
    if not nonempty:
        return False
    prefixed = [
        value
        for value in nonempty
        if re.search(r"\d{4,}$", _trim(value)) and not re.fullmatch(r"\d+", _trim(value))
    ]
    return len(prefixed) / len(nonempty) >= 0.6


def _value_set(values: list[Any], normalizer) -> set[str]:
    return {normalizer(value) for value in values if not _blank(value) and normalizer(value)}


def _overlap(left: set[str], right: set[str]) -> tuple[int, float, float]:
    if not left or not right:
        return 0, 0.0, 0.0
    count = len(left & right)
    return count, count / len(left), count / len(right)


def _join_candidate(left: dict, right: dict, left_col: str, right_col: str) -> dict | None:
    left_values = _source_values(left, left_col)
    right_values = _source_values(right, right_col)
    strategies = [
        ("exact", _norm_exact, None),
        ("alphanumeric", _norm_alnum, "remove punctuation/case before joining"),
        ("trailing_digits", _norm_suffix_digits, "extract trailing digits from the fuller key before joining"),
    ]
    best: dict[str, Any] | None = None
    for strategy, normalizer, transform_hint in strategies:
        left_set = _value_set(left_values, normalizer)
        right_set = _value_set(right_values, normalizer)
        count, left_ratio, right_ratio = _overlap(left_set, right_set)
        if count == 0:
            continue
        score = min(left_ratio, right_ratio)
        candidate = {
            "leftColumn": left_col,
            "rightColumn": right_col,
            "matchStrategy": strategy,
            "overlapCount": count,
            "overlapRatioLeft": round(left_ratio, 4),
            "overlapRatioRight": round(right_ratio, 4),
            "uniqueRatioLeft": round(_unique_ratio(left, left_col), 4),
            "uniqueRatioRight": round(_unique_ratio(right, right_col), 4),
            "transformHint": transform_hint,
            "confidence": "high" if score >= 0.8 else ("medium" if score >= 0.25 else "low"),
        }
        if best is None or min(candidate["overlapRatioLeft"], candidate["overlapRatioRight"]) > min(best["overlapRatioLeft"], best["overlapRatioRight"]):
            best = candidate
    if best and (best["confidence"] != "low" or KEY_NAME_RE.search(left_col) or KEY_NAME_RE.search(right_col)):
        return best
    return None


def _derived_key_hint(left: dict, right: dict, candidate: dict[str, Any]) -> dict[str, Any] | None:
    if candidate.get("matchStrategy") != "trailing_digits":
        return None
    left_col = str(candidate.get("leftColumn") or "")
    right_col = str(candidate.get("rightColumn") or "")
    if not left_col or not right_col:
        return None
    left_values = _source_values(left, left_col)
    right_values = _source_values(right, right_col)
    derived_steps = []
    left_join_field = left_col
    right_join_field = right_col
    if _has_prefixed_trailing_digits(left_values):
        left_join_field = f"{left_col} Digits"
        derived_steps.append(
            {
                "source": "left",
                "sourceName": left.get("sourceName"),
                "sourceColumn": left_col,
                "derivedColumn": left_join_field,
                "builderHint": f'nb.op_expr("{left_join_field}", \'REGEX_EXTRACT(`{left_col}`, "([0-9]+)$")\', "string")',
            }
        )
    if _has_prefixed_trailing_digits(right_values):
        right_join_field = f"{right_col} Digits"
        derived_steps.append(
            {
                "source": "right",
                "sourceName": right.get("sourceName"),
                "sourceColumn": right_col,
                "derivedColumn": right_join_field,
                "builderHint": f'nb.op_expr("{right_join_field}", \'REGEX_EXTRACT(`{right_col}`, "([0-9]+)$")\', "string")',
            }
        )
    if not derived_steps:
        return None
    if _mostly_digits(left_values) and left_join_field == left_col:
        left_join_field = left_col
    if _mostly_digits(right_values) and right_join_field == right_col:
        right_join_field = right_col
    return {
        "kind": "trailing_digits_derived_key",
        "leftColumn": left_col,
        "rightColumn": right_col,
        "leftJoinField": left_join_field,
        "rightJoinField": right_join_field,
        "overlapCount": candidate.get("overlapCount"),
        "overlapRatioLeft": candidate.get("overlapRatioLeft"),
        "overlapRatioRight": candidate.get("overlapRatioRight"),
        "requiresConfirmation": True,
        "reason": "Trailing digits from one identifier overlap the other identifier; confirm before Builder uses this derived key.",
        "derivedSteps": derived_steps,
    }


def _is_conservative_join(candidate: dict[str, Any]) -> bool:
    """True only for join evidence safe enough to use without AI review."""
    left_col = str(candidate.get("leftColumn") or "")
    right_col = str(candidate.get("rightColumn") or "")
    if candidate.get("confidence") != "high":
        return False
    if candidate.get("matchStrategy") not in {"exact", "alphanumeric"}:
        return False
    if int(candidate.get("overlapCount") or 0) < 2:
        return False
    if min(float(candidate.get("overlapRatioLeft") or 0), float(candidate.get("overlapRatioRight") or 0)) < 0.95:
        return False
    if max(float(candidate.get("uniqueRatioLeft") or 0), float(candidate.get("uniqueRatioRight") or 0)) < 0.95:
        return False
    if _column_norm(left_col) == _column_norm(right_col):
        return bool(IDENTIFIER_NAME_RE.search(left_col) or IDENTIFIER_NAME_RE.search(right_col))
    return bool(IDENTIFIER_NAME_RE.search(left_col) and IDENTIFIER_NAME_RE.search(right_col))


def _is_ai_review_join_candidate(candidate: dict[str, Any]) -> bool:
    left_col = str(candidate.get("leftColumn") or "")
    right_col = str(candidate.get("rightColumn") or "")
    if int(candidate.get("overlapCount") or 0) < 2:
        return False
    if max(float(candidate.get("uniqueRatioLeft") or 0), float(candidate.get("uniqueRatioRight") or 0)) < 0.5:
        return False
    return bool(
        IDENTIFIER_NAME_RE.search(left_col)
        or IDENTIFIER_NAME_RE.search(right_col)
        or KEY_NAME_RE.search(left_col)
        or KEY_NAME_RE.search(right_col)
    )


def _column_norm(name: str) -> str:
    return re.sub(r"[^a-z0-9]", "", name.lower())


def _candidate_join_pairs(left: dict, right: dict) -> list[tuple[str, str]]:
    pairs: set[tuple[str, str]] = set()
    left_keys = [item["column"] for item in left["keyCandidates"]]
    right_keys = [item["column"] for item in right["keyCandidates"]]
    for left_col in left_keys:
        for right_col in right_keys:
            pairs.add((left_col, right_col))

    right_by_norm: dict[str, list[str]] = defaultdict(list)
    for column in right["_columns"]:
        if KEY_SUPPRESS_RE.search(column):
            continue
        right_by_norm[_column_norm(column)].append(column)
    for left_col in left["_columns"]:
        if KEY_SUPPRESS_RE.search(left_col):
            continue
        for right_col in right_by_norm.get(_column_norm(left_col), []):
            pairs.add((left_col, right_col))
    return sorted(pairs)


def relationship_profiles(sources: list[dict], *, max_candidates_per_pair: int = 8) -> list[dict]:
    relationships = []
    tabular_sources = [source for source in sources if source.get("sourceKind") == "tabular"]
    for left_index, left in enumerate(tabular_sources):
        for right in tabular_sources[left_index + 1:]:
            suggested = []
            ai_review = []
            derived_hints = []
            for left_col, right_col in _candidate_join_pairs(left, right):
                candidate = _join_candidate(left, right, left_col, right_col)
                if candidate:
                    if _is_conservative_join(candidate):
                        suggested.append(candidate)
                    elif _is_ai_review_join_candidate(candidate):
                        ai_review.append(candidate)
                    hint = _derived_key_hint(left, right, candidate)
                    if hint:
                        derived_hints.append(hint)
            suggested = sorted(
                suggested,
                key=lambda item: (
                    -min(item["overlapRatioLeft"], item["overlapRatioRight"]),
                    -item["overlapCount"],
                    item["leftColumn"],
                    item["rightColumn"],
                ),
            )[:max_candidates_per_pair]
            ai_review = sorted(
                ai_review,
                key=lambda item: (
                    -min(item["overlapRatioLeft"], item["overlapRatioRight"]),
                    -item["overlapCount"],
                    item["leftColumn"],
                    item["rightColumn"],
                ),
            )[:max_candidates_per_pair]
            derived_hints = sorted(
                derived_hints,
                key=lambda item: (
                    -min(float(item.get("overlapRatioLeft") or 0), float(item.get("overlapRatioRight") or 0)),
                    -int(item.get("overlapCount") or 0),
                    item["leftColumn"],
                    item["rightColumn"],
                ),
            )[:max_candidates_per_pair]
            relationships.append(
                {
                    "leftSource": left["sourceName"],
                    "rightSource": right["sourceName"],
                    "suggestedJoins": suggested,
                    "aiReviewJoinCandidates": ai_review,
                    "derivedKeyHints": derived_hints,
                    "warnings": [] if suggested else ["No conservative join suggestion; AI review is required before using overlaps as join keys."],
                }
            )
    return relationships


def _profile_source(
    *,
    source_name: str,
    path: Path,
    connector: str,
    columns: list[str],
    rows: list[list[Any]],
    truncated: bool,
    sheet: str | None,
    header_row: int,
    hint_applied: dict[str, Any] | None,
    sample_limit: int,
) -> dict[str, Any]:
    observed_schema, date_candidates = column_profiles(columns, rows, sample_limit=sample_limit)
    row_count = len(rows)
    warnings = []
    embedded_headers = _embedded_header_rows(rows, header_row=header_row)
    if truncated:
        warnings.append("Profile scan hit the row limit; row counts and distinct counts are based on scanned rows.")
    if header_row > 1:
        warnings.append(f"Detected source header on row {header_row}; leading rows were treated as preamble.")
    if embedded_headers:
        warnings.append(
            "Possible additional table header row(s) inside the profiled range: "
            + ", ".join(str(row) for row in embedded_headers[:5])
            + ("." if len(embedded_headers) <= 5 else ", ...")
        )
    if date_candidates:
        needs_adapter = [item["column"] for item in date_candidates if item.get("needsAdapter")]
        if needs_adapter:
            warnings.append("Date-like fields need a source adapter or transform before business logic: " + ", ".join(needs_adapter))
    profile = {
        "sourceName": source_name,
        "sourceKind": "tabular",
        "path": str(path),
        "sheet": sheet,
        "headerRow": header_row,
        "hintApplied": hint_applied or {},
        "connector": connector,
        "rowCount": row_count,
        "rowCountIsCapped": truncated,
        "profileConfidence": "hinted" if hint_applied else ("needs_ai_review" if embedded_headers else ("inferred" if header_row > 1 else "deterministic")),
        "needsAiProfile": bool(embedded_headers),
        "embeddedHeaderRows": embedded_headers,
        "columnCount": len(columns),
        "observed_schema": observed_schema,
        "expected_schema": expected_schema_from_profile(observed_schema, date_candidates),
        "dateCandidates": date_candidates,
        "keyCandidates": key_candidates(observed_schema, row_count),
        "warnings": warnings,
        "sourcePlanSuggestion": {
            "source_name": source_name,
            "business_role": "",
            "source_kind": "tabular",
            "dataset": {
                "dataset_id": "",
                "profile_status": "full_profile",
                "connector": connector,
                "row_count": row_count,
                "observed_schema": observed_schema,
            },
            "expected_schema": expected_schema_from_profile(observed_schema, date_candidates),
        },
        "_columns": columns,
        "_rows": rows,
    }
    return profile


def _profile_pdf(path: Path) -> dict[str, Any]:
    observed_schema = [
        {
            "name": "Content",
            "dataType": "binary",
            "sampleValues": [path.name],
            "nonEmptyCount": 1,
            "distinctCount": 1,
        }
    ]
    return {
        "sourceName": path.stem,
        "sourceKind": "binary_document",
        "path": str(path),
        "sheet": None,
        "connector": "pdf",
        "rowCount": 1,
        "rowCountIsCapped": False,
        "columnCount": 1,
        "observed_schema": observed_schema,
        "expected_schema": [],
        "dateCandidates": [],
        "keyCandidates": [],
        "warnings": ["PDF content is binary; profile only confirms the file source, not extracted fields."],
        "sourcePlanSuggestion": {
            "source_name": path.stem,
            "business_role": "",
            "source_kind": "binary_document",
            "dataset": {
                "dataset_id": "",
                "profile_status": "full_profile",
                "connector": "pdf",
                "row_count": 1,
                "observed_schema": observed_schema,
            },
            "extraction": {
                "prompt": "",
            },
        },
        "_columns": [],
        "_rows": [],
    }


def ai_profile_requests(public_sources: list[dict[str, Any]]) -> list[dict[str, Any]]:
    requests = []
    for source in public_sources:
        reasons = []
        if source.get("needsAiProfile"):
            reasons.append("Profiler detected possible additional table headers inside the scanned range.")
        if source.get("profileConfidence") == "inferred":
            reasons.append("Profiler inferred a non-first header row; AI can confirm or override it.")
        if not reasons:
            continue
        requests.append(
            {
                "sourceName": source.get("sourceName"),
                "path": source.get("path"),
                "sheet": source.get("sheet"),
                "reasons": reasons,
                "hintTemplate": {
                    "path": source.get("path"),
                    "sources": [
                        {
                            "sheet": source.get("sheet"),
                            "sourceName": source.get("sourceName"),
                            "headerRow": source.get("headerRow"),
                            "dataStartRow": (source.get("headerRow") or 1) + 1,
                            "dataEndRow": None,
                            "note": "AI-confirmed table range; split into multiple entries when the sheet contains multiple tables.",
                        }
                    ],
                },
            }
        )
    return requests


def profile_files(
    files: list[Path],
    *,
    sheets: list[str] | None = None,
    hints: list[dict[str, Any]] | None = None,
    delimiter: str | list[str] = ",",
    max_rows: int = DEFAULT_SCAN_ROWS,
    sample_limit: int = DEFAULT_SAMPLE_VALUES,
) -> dict[str, Any]:
    sources: list[dict[str, Any]] = []
    errors: list[str] = []
    warnings: list[str] = []
    hints = hints or []
    # One delimiter per --file, in order; the last one given applies to any further files.
    delimiters = [delimiter] if isinstance(delimiter, str) else (list(delimiter) or [","])
    for file_index, raw_path in enumerate(files):
        file_delimiter = delimiters[min(file_index, len(delimiters) - 1)]
        path = raw_path.expanduser().resolve()
        if not path.exists() or not path.is_file():
            errors.append(f"File does not exist: {raw_path}")
            continue
        suffix = path.suffix.lower()
        file_hints = _hints_for_file(hints, path)
        try:
            if suffix == ".pdf":
                sources.append(_profile_pdf(path))
            elif suffix in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
                if file_hints:
                    source_specs = file_hints
                else:
                    source_specs = [{"sheet": sheet, "_auto": True} for sheet in list(sheets or excel_sheet_names(path))]
                for spec in source_specs:
                    sheet = str(spec.get("sheet") or spec.get("tab") or "").strip()
                    if not sheet:
                        errors.append(f"{path}: Excel source hint is missing `sheet`.")
                        continue
                    columns, rows, truncated, header_row = _read_xlsx(
                        path,
                        sheet,
                        max_rows=max_rows,
                        header_row=_hint_int(spec, "headerRow", "header_row"),
                        data_start_row=_hint_int(spec, "dataStartRow", "data_start_row"),
                        data_end_row=_hint_int(spec, "dataEndRow", "data_end_row"),
                    )
                    source_name = str(spec.get("sourceName") or spec.get("source_name") or "").strip()
                    if not source_name:
                        source_name = f"{path.stem} - {sheet}" if len(source_specs) > 1 else path.stem
                    hint_applied = {
                        key: spec[key]
                        for key in ("sheet", "sourceName", "source_name", "headerRow", "header_row", "dataStartRow", "data_start_row", "dataEndRow", "data_end_row", "note")
                        if key in spec
                    }
                    if spec.get("_auto"):
                        hint_applied = {}
                    sources.append(
                        _profile_source(
                            source_name=source_name,
                            path=path,
                            connector="excel",
                            columns=columns,
                            rows=rows,
                            truncated=truncated,
                            sheet=sheet,
                            header_row=header_row,
                            hint_applied=hint_applied,
                            sample_limit=sample_limit,
                        )
                    )
            elif suffix in {".csv", ".tsv", ".txt"}:
                effective_delimiter = "\t" if suffix == ".tsv" and file_delimiter == "," else file_delimiter
                columns, rows, truncated, encoding = _read_csv_with_encoding(path, delimiter=effective_delimiter, max_rows=max_rows)
                if encoding != "utf-8":
                    warnings.append(f"{path.name}: not valid UTF-8; read as {encoding}. Set the Savant dataset charset to match the file (the upload defaults to UTF-8).")
                if len(columns) == 1 and effective_delimiter not in columns[0] and any(sep in columns[0] for sep in (";", ",", "\t", "|")):
                    warnings.append(f"{path.name}: only one column parsed with delimiter {effective_delimiter!r}; the header looks like it uses a different separator — pass --delimiter per file, in --file order.")
                sources.append(
                    _profile_source(
                        source_name=str((file_hints[0] if file_hints else {}).get("sourceName") or (file_hints[0] if file_hints else {}).get("source_name") or path.stem),
                        path=path,
                        connector="csv",
                        columns=columns,
                        rows=rows,
                        truncated=truncated,
                        sheet=None,
                        header_row=1,
                        hint_applied=file_hints[0] if file_hints else None,
                        sample_limit=sample_limit,
                    )
                )
            elif suffix == ".xls":
                errors.append(f"{path}: legacy .xls files are not supported by the local profiler; convert to .xlsx or CSV.")
            else:
                errors.append(f"{path}: unsupported source file type `{suffix or '(none)'}`; use CSV, TSV, XLSX/XLSM, or PDF.")
        except Exception as exc:  # noqa: BLE001
            errors.append(f"{path}: {type(exc).__name__}: {exc}")
    public_sources = []
    for source in sources:
        public_sources.append({key: value for key, value in source.items() if not key.startswith("_")})
    ai_requests = ai_profile_requests(public_sources)
    ready_for_source_planning = bool(public_sources) and not errors and not ai_requests
    status = "ready" if ready_for_source_planning else ("needs_ai_profile" if public_sources and not errors else ("partial" if public_sources else "blocked"))
    report = {
        "task": "source_profile",
        "profileVersion": 1,
        "scan": {
            "rowLimit": max_rows,
            "sampleValueLimit": sample_limit,
            "hintsApplied": bool(hints),
        },
        "sources": public_sources,
        "relationships": relationship_profiles(sources),
        "aiProfileRequests": ai_requests,
        "readyForSourcePlanning": ready_for_source_planning,
        "errors": errors,
        "warnings": warnings,
        "status": status,
    }
    report["brief"] = brief_report(report)
    return report


def _source_location(source: dict[str, Any]) -> str:
    location = str(source.get("path") or "")
    if source.get("sheet"):
        location += f"#{source['sheet']}"
    return location


def _brief_columns(source: dict[str, Any]) -> list[str]:
    columns = []
    for column in source.get("observed_schema", []):
        if isinstance(column, dict):
            columns.append(f"{column.get('name')}:{column.get('dataType')}")
    return columns


def _brief_date_adapters(source: dict[str, Any]) -> list[dict[str, Any]]:
    return [
        {
            "column": item.get("column"),
            "expectedDataType": item.get("expectedDataType"),
            "reason": item.get("reason"),
        }
        for item in source.get("dateCandidates", [])
        if isinstance(item, dict) and item.get("needsAdapter")
    ]


def _brief_join(join: dict[str, Any]) -> str:
    left = join.get("leftColumn")
    right = join.get("rightColumn")
    strategy = join.get("matchStrategy")
    overlap = min(float(join.get("overlapRatioLeft") or 0), float(join.get("overlapRatioRight") or 0))
    return f"{left} -> {right} ({strategy}, overlap {overlap:.0%})"


def _brief_derived_key_hint(hint: dict[str, Any]) -> str:
    left = hint.get("leftColumn")
    right = hint.get("rightColumn")
    left_join = hint.get("leftJoinField")
    right_join = hint.get("rightJoinField")
    overlap = min(float(hint.get("overlapRatioLeft") or 0), float(hint.get("overlapRatioRight") or 0))
    return (
        f"{left} -> {right}: derive trailing digits, then join "
        f"{left_join} -> {right_join} (overlap {overlap:.0%}, requires confirmation)"
    )


def brief_report(report: dict[str, Any]) -> dict[str, Any]:
    """Compact, first-read summary for Planner/Builder agents.

    The full report remains the evidence source. This brief intentionally omits
    samples and full schema metadata so agents can scan readiness and only open
    the detailed sections when needed.
    """
    sources = report.get("sources") or []
    relationships = report.get("relationships") or []
    ai_requests = report.get("aiProfileRequests") or []
    return {
        "task": report.get("task"),
        "status": report.get("status"),
        "readyForSourcePlanning": report.get("readyForSourcePlanning"),
        "sourceCount": len(sources),
        "sources": [
            {
                "sourceName": source.get("sourceName"),
                "sourceKind": source.get("sourceKind"),
                "location": _source_location(source),
                "rowCount": source.get("rowCount"),
                "rowCountIsCapped": source.get("rowCountIsCapped"),
                "profileConfidence": source.get("profileConfidence"),
                "columns": _brief_columns(source),
                "dateAdaptersNeeded": _brief_date_adapters(source),
                "keyCandidates": [item.get("column") for item in source.get("keyCandidates", []) if isinstance(item, dict)],
                "warnings": source.get("warnings") or [],
            }
            for source in sources
            if isinstance(source, dict)
        ],
        "relationships": [
            {
                "leftSource": relationship.get("leftSource"),
                "rightSource": relationship.get("rightSource"),
                "suggestedJoins": [_brief_join(join) for join in relationship.get("suggestedJoins", []) if isinstance(join, dict)],
                "aiReviewJoinCandidates": [_brief_join(join) for join in relationship.get("aiReviewJoinCandidates", []) if isinstance(join, dict)],
                "derivedKeyHints": [
                    _brief_derived_key_hint(hint)
                    for hint in relationship.get("derivedKeyHints", [])
                    if isinstance(hint, dict)
                ],
                "warnings": relationship.get("warnings") or [],
            }
            for relationship in relationships
            if isinstance(relationship, dict)
        ],
        "aiProfileRequests": [
            {
                "sourceName": request.get("sourceName"),
                "location": _source_location(request),
                "reasons": request.get("reasons") or [],
            }
            for request in ai_requests
            if isinstance(request, dict)
        ],
        "errors": report.get("errors") or [],
    }


def default_output_path() -> Path:
    return workspace_tmp("source-profiles", "source-profile.json")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--file", action="append", required=True, type=Path, help="Local CSV/TSV/XLSX/PDF file to profile. Repeatable.")
    parser.add_argument("--sheet", action="append", help="Excel sheet name to profile. Repeatable; applies to Excel files.")
    parser.add_argument("--hints-json", type=Path, help="Optional AI-supplied source/table hints JSON from a prior ambiguous profile.")
    parser.add_argument("--delimiter", action="append", help="CSV delimiter. Repeatable, one per --file in order; the last one covers remaining files. Defaults to comma.")
    parser.add_argument("--max-rows", type=int, default=DEFAULT_SCAN_ROWS, help="Maximum data rows to scan per source.")
    parser.add_argument("--sample-values", type=int, default=DEFAULT_SAMPLE_VALUES, help="Distinct sample values per column.")
    parser.add_argument("--output-path", type=Path, help="Where to write the JSON report.")
    parser.add_argument("--brief-output-path", type=Path, help="Optional path for a compact brief JSON report.")
    parser.add_argument("--stdout", action="store_true", help="Print the JSON report to stdout after writing it.")
    parser.add_argument("--stdout-format", choices=("full", "brief"), default="full", help="When --stdout is set, choose full report JSON or compact brief JSON.")
    args = parser.parse_args(argv)

    if args.max_rows <= 0:
        parser.error("--max-rows must be positive.")
    if args.sample_values <= 0:
        parser.error("--sample-values must be positive.")
    try:
        hints = _load_hints(args.hints_json)
        report = profile_files(
            args.file,
            sheets=args.sheet,
            hints=hints,
            delimiter=args.delimiter or [","],
            max_rows=args.max_rows,
            sample_limit=args.sample_values,
        )
    except Exception as exc:  # noqa: BLE001
        print(f"source profile: {type(exc).__name__}: {exc}", file=sys.stderr)
        return 2
    output = args.output_path or default_output_path()
    save_json(report, output)
    if args.brief_output_path:
        save_json(report["brief"], args.brief_output_path)
    if args.stdout:
        payload = report["brief"] if args.stdout_format == "brief" else report
        print(json.dumps(payload, indent=2))
    else:
        source_count = len(report["sources"])
        relationship_count = sum(1 for rel in report["relationships"] if rel.get("suggestedJoins"))
        print(f"Profiled {source_count} source(s); found join candidates for {relationship_count} pair(s).")
        print(f"Wrote source profile to {output}")
        for warning in report.get("warnings", []):
            print(f"warning: {warning}", file=sys.stderr)
        for error in report.get("errors", []):
            print(f"error: {error}", file=sys.stderr)
        if args.brief_output_path:
            print(f"Wrote source profile brief to {args.brief_output_path}")
    return 0 if report["status"] != "blocked" else 2


if __name__ == "__main__":
    raise SystemExit(main())
