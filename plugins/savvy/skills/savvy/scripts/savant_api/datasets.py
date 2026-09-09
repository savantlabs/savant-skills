from __future__ import annotations

import argparse
import csv
import datetime as _dt
import json
import time
from pathlib import Path

from . import sources
from . import cli as api
from .fileio import save_json, workspace_tmp
from .recipe_input import load_records
from .models import SavantAppApiError, SavantSessionContext

OPENPYXL_DEPENDENCY = "openpyxl>=3.1,<4"


def _load_openpyxl():
    try:
        import openpyxl
    except ImportError as exc:
        raise SavantAppApiError(f"Excel dataset creation requires Python package `{OPENPYXL_DEPENDENCY}`.") from exc
    return openpyxl


def infer_type(values: list[str]) -> str:
    """integer / number / string from sampled values (matches the upload wizard's behavior)."""
    nonempty = [v for v in values if v is not None and str(v).strip() != ""]
    if not nonempty:
        return "string"
    def is_int(v):
        try: int(str(v)); return True
        except ValueError: return False
    def is_float(v):
        try: float(str(v)); return True
        except ValueError: return False
    if all(is_int(v) for v in nonempty):
        return "integer"
    if all(is_float(v) for v in nonempty):
        return "number"
    return "string"


def coerce(value, dtype: str):
    if value is None or value == "":
        return None
    if dtype == "integer":
        try: return int(value)
        except (ValueError, TypeError): return value
    if dtype == "number":
        try: return float(value)
        except (ValueError, TypeError): return value
    return value


def parse_csv(path: Path, delimiter: str, sample_rows: int):
    with open(path, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.reader(f, delimiter=delimiter))
    if not rows:
        raise SavantAppApiError(f"file has no rows: {path}")
    header, body = rows[0], rows[1:]
    cols = list(zip(*body)) if body else [[] for _ in header]
    types = [infer_type(list(c)) for c in cols] if body else ["string"] * len(header)
    schema = [{"id": h, "name": h, "dataType": t} for h, t in zip(header, types)]
    data = [[coerce(v, schema[i]["dataType"]) for i, v in enumerate(r)] for r in body[:sample_rows]]
    return header, schema, data


def build_create_body(name: str, file_id: str, header, schema, data, delimiter: str, length: int = 0) -> dict:
    config = {
        "fileType": "CSV", "overwrittenTypes": {}, "selected": list(header),
        "charset": "UTF_8", "catalog": "catalog", "schema": "schema", "table": "file",
        "delimiter": delimiter, "qualifier": '"', "escape": "\\",
        "readAsStored": True, "extractionMethod": "native", "autoAppendNewFields": False,
        "fullSample": {"length": int(length), "schema": schema, "data": data,
                       "createdTime": int(time.time() * 1000), "hasDataAfterSkipRows": True},
        "sample": {"schema": schema, "data": data},
        "type": "file", "connector": "csv", "dataFormat": "CSV",
        "fileId": file_id, "tabName": "file",
        "fileParserProps": {"delimiter": delimiter, "qualifier": '"', "escape": "\\", "charset": "UTF_8"},
        "passwordProtected": False,
    }
    return {"name": name, "config": config, "profile": {"schema": schema, "sample": data}}


def _xlsx_type(values) -> str:
    nonempty = [v for v in values if v is not None and str(v).strip() != ""]
    if not nonempty:
        return "string"
    if all(isinstance(v, bool) for v in nonempty):
        return "boolean"
    if all(isinstance(v, int) and not isinstance(v, bool) for v in nonempty):
        return "integer"
    if all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in nonempty):
        return "number"
    if all(isinstance(v, _dt.datetime) for v in nonempty):
        return "datetime"
    if all(isinstance(v, _dt.date) for v in nonempty):
        return "date"
    return "string"


def _json_safe(v):
    if isinstance(v, (_dt.datetime, _dt.date)):
        return v.isoformat()
    return v


def _excel_stored_value(cell, epoch):
    """Return the value Savant's Excel "Read As Stored" mode exposes for a cell.

    openpyxl helpfully converts date-formatted Excel serial numbers into Python date/datetime
    objects. Savant's Read As Stored mode keeps the underlying serial number, so convert those
    values back before inferring schema/sample types.
    """
    value = cell.value
    if value is None:
        return None
    if getattr(cell, "is_date", False) and isinstance(value, (_dt.datetime, _dt.date)):
        from openpyxl.utils.datetime import to_excel
        serial = to_excel(value, epoch)
        return int(serial) if float(serial).is_integer() else serial
    return value


def match_sheets(sheet_names: list[str], pattern: str) -> list[str]:
    import fnmatch
    matched = [s for s in sheet_names if fnmatch.fnmatch(s, pattern)]
    if not matched:
        raise SavantAppApiError(f"no sheets match pattern {pattern!r}; sheets: {sheet_names}")
    return matched


def parse_xlsx(path: Path, sheet: str | None, sample_rows: int, skip_rows: int = 0):
    openpyxl = _load_openpyxl()
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet_names = list(wb.sheetnames)
    ws = wb[sheet] if sheet else wb[sheet_names[0]]
    rows = [[_excel_stored_value(cell, wb.epoch) for cell in row] for row in ws.iter_rows()]
    if not rows:
        raise SavantAppApiError(f"sheet has no rows: {path}")
    if skip_rows < 0:
        raise SavantAppApiError("skip_rows must be zero or greater.")
    if skip_rows >= len(rows):
        raise SavantAppApiError(f"skip_rows={skip_rows} skips all rows in sheet: {path}#{ws.title}")
    header = [str(c) if c is not None else f"col_{i}" for i, c in enumerate(rows[skip_rows])]
    body = rows[skip_rows + 1:]
    cols = list(zip(*body)) if body else [[] for _ in header]
    types = [_xlsx_type(list(c)) for c in cols] if body else ["string"] * len(header)
    schema = [{"id": h, "name": h, "dataType": t} for h, t in zip(header, types)]
    data = [[_json_safe(v) for v in r] for r in body[:sample_rows]]
    return ws.title, sheet_names, header, schema, data


def build_excel_body(name, file_id, sheet, sheet_names, header, schema, data,
                     use_tab_pattern: bool = False, tab_pattern: str = "*", skip_rows: int = 0,
                     length: int = 0) -> dict:
    config = {
        "fileType": "EXCEL", "overwrittenTypes": {}, "selected": list(header),
        "charset": "UTF_8", "catalog": "catalog", "schema": "schema", "table": sheet,
        "readAsStored": True, "extractionMethod": "native", "autoAppendNewFields": False,
        "fullSample": {"length": int(length), "schema": schema, "data": data,
                       "createdTime": int(time.time() * 1000), "hasDataAfterSkipRows": True},
        "sample": {"schema": schema, "data": data},
        "type": "file", "connector": "excel", "dataFormat": "EXCEL",
        "fileId": file_id, "tabName": sheet, "tabNames": list(sheet_names),
        "useTabPattern": bool(use_tab_pattern), "tabPattern": tab_pattern, "passwordProtected": False,
        "skipRows": int(skip_rows),
    }
    return {"name": name, "config": config, "profile": {"schema": schema, "sample": data}}


def build_binary_body(name, file_id, file_name: str, file_type: str, connector: str, data_format: str) -> dict:
    """PDF / image / blob: a single binary `Content` column whose sample value is the file name.
    The actual content is extracted downstream by a vision step; the dataset just references the
    uploaded file. (Sample must be non-empty — create-async rejects an empty sample.)"""
    schema = [{"id": "content", "name": "Content", "dataType": "binary"}]
    data = [[file_name]]
    config = {
        "fileType": file_type, "overwrittenTypes": {}, "selected": ["Content"],
        "charset": "UTF_8", "catalog": "catalog", "schema": "schema", "table": "file",
        "readAsStored": True, "extractionMethod": "native", "autoAppendNewFields": False,
        "fullSample": {"length": 0, "schema": schema, "data": data,
                       "createdTime": int(time.time() * 1000), "hasDataAfterSkipRows": True},
        "sample": {"schema": schema, "data": data},
        "type": "file", "connector": connector, "dataFormat": data_format,
        "fileId": file_id, "tabName": "file", "passwordProtected": False,
    }
    return {"name": name, "config": config, "profile": {"schema": schema, "sample": data}}


_CONTENT_TYPE = {
    "csv": "text/csv",
    "excel": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}


def detect_format(path: Path) -> str:
    ext = path.suffix.lower().lstrip(".")
    if ext in ("xlsx", "xls", "xlsm", "xlsb"):
        return "excel"
    if ext == "pdf":
        return "pdf"
    return "csv"


def _has_ci_duplicate_names(names) -> bool:
    """True when ``names`` contains a repeat under the server's case-insensitive comparison."""
    lowered = [str(n).strip().lower() for n in names]
    return len(set(lowered)) != len(lowered)


def _visible_excel_sheets(path: Path) -> list[str]:
    """Sheet names the upload wizard would offer — visible tabs only. SAP BI exports carry a
    very-hidden control sheet (``_com.sap.ip.bi.xl.hiddensheet``); excluding hidden/very-hidden
    sheets keeps the dataset's tab list (and the UI's Tab dropdown default) on real data tabs."""
    openpyxl = _load_openpyxl()
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        return [ws.title for ws in wb.worksheets if getattr(ws, "sheet_state", "visible") == "visible"]
    finally:
        wb.close()


def _parts_from_server_sample(sample: dict):
    """Turn a server profile (``sources.sample_uploaded_file``) into
    (header, schema, data, length) carrying the server's deduplicated names and inferred types —
    the same source of truth the upload wizard builds its dataset from."""
    fields = sample.get("schema") or []
    header = [str(f.get("name")) for f in fields]
    schema = [
        {"id": f.get("id") or f.get("name"), "name": f.get("name"),
         "dataType": f.get("dataType") or "string"}
        for f in fields
    ]
    return header, schema, sample.get("data") or [], int(sample.get("length") or 0)


def _dedupe_names_like_server(header, schema):
    """Best-effort local fallback that approximates the server's dedupe (first occurrence kept;
    later case-insensitive repeats get ``_2``/``_3``; blank headers become ``Column``). Only used
    when the server profile call fails, so the offline path still produces unique column names."""
    seen: dict[str, int] = {}
    names: list[str] = []
    for raw in header:
        base = str(raw).strip() or "Column"
        key = base.lower()
        seen[key] = seen.get(key, 0) + 1
        names.append(base if seen[key] == 1 else f"{base}_{seen[key]}")
    new_schema = [
        {"id": names[i], "name": names[i], "dataType": schema[i]["dataType"]}
        for i in range(len(names))
    ]
    return names, new_schema


def create_dataset(context, path: Path, name: str, fmt: str = "auto",
                   delimiter: str = ",", sheet: str | None = None, tab_pattern: str | None = None,
                   sample_rows: int = 100, skip_rows: int = 0) -> dict:
    """Upload + create a dataset of the given format (auto-detected by extension if 'auto').
    For Excel: pass `sheet` for a single tab, or `tab_pattern` (e.g. "*" or "2024-*") to combine
    multiple tabs. Returns the created source object (result.id = new dataset id).

    Uses the raw API primitives in `savant_api.sources` (`upload_file_async`, `create_source`);
    this layer only owns file-format parsing and building the create body."""
    if fmt == "auto":
        fmt = detect_format(path)
    file_id = sources.upload_file_async(context, path, _CONTENT_TYPE.get(fmt, "application/octet-stream"))
    if fmt == "csv":
        # Prefer the server's profile (deduped names, inferred types, ~1000-row sample) so an
        # API-created dataset matches a manual upload; fall back to a local parse if it fails.
        try:
            sample = sources.sample_uploaded_file(context, file_id, data_format="CSV", delimiter=delimiter)
            header, schema, data, length = _parts_from_server_sample(sample)
        except SavantAppApiError:
            header, schema, data = parse_csv(path, delimiter, sample_rows)
            length = 0
            if _has_ci_duplicate_names(header):
                header, schema = _dedupe_names_like_server(header, schema)
        body = build_create_body(name, file_id, header, schema, data, delimiter, length=length)
    elif fmt == "excel":
        visible = _visible_excel_sheets(path)
        if tab_pattern:
            # multi-tab: schema/sample come from the first matching sheet (tabs are stacked by name)
            openpyxl = _load_openpyxl()
            all_names = openpyxl.load_workbook(path, read_only=True).sheetnames
            tab = match_sheets(all_names, tab_pattern)[0]
        else:
            tab = sheet or (visible[0] if visible else None)
        tab_names = visible if visible else None
        if tab_names is not None and tab is not None and tab not in tab_names:
            tab_names = [tab, *tab_names]
        try:
            sample = sources.sample_uploaded_file(
                context, file_id, data_format="EXCEL", tab=tab, skip_rows=skip_rows)
            header, schema, data, length = _parts_from_server_sample(sample)
        except SavantAppApiError:
            sht, sheet_names, header, schema, data = parse_xlsx(path, tab, sample_rows, skip_rows=skip_rows)
            tab, length = sht, 0
            if tab_names is None:
                tab_names = sheet_names
            if _has_ci_duplicate_names(header):
                header, schema = _dedupe_names_like_server(header, schema)
        body = build_excel_body(
            name, file_id, tab, tab_names or [], header, schema, data,
            use_tab_pattern=bool(tab_pattern), tab_pattern=(tab_pattern or "*"),
            skip_rows=skip_rows, length=length)
    elif fmt == "pdf":
        # Send fileType/connector "PDF"/"pdf"; the server stores it as a binary dataset
        # (connector "binary", dataFormat "BINARY"). Sending "binary" directly is rejected.
        body = build_binary_body(name, file_id, path.name, "PDF", "pdf", "PDF")
    else:
        raise SavantAppApiError(f"unsupported format: {fmt}")
    return sources.create_source(context, body)


def create_csv_dataset(context, path: Path, name: str, delimiter: str = ",", sample_rows: int = 100) -> dict:
    """Back-compat wrapper for CSV creation."""
    return create_dataset(context, path, name, "csv", delimiter=delimiter, sample_rows=sample_rows)


def create_datasets_bulk(context, manifest: list[dict], *, existing_json: Path | None = None) -> dict:
    """Create several datasets from one manifest — the deterministic loop the AI should not hand-run.

    Manifest items: {"file": <path>, "name": <display name>, and optional "type", "sheet",
    "tab_pattern", "skip_rows", "delimiter", "sample_rows"}.

    Behavior contracts:
    - **Idempotent by name**: an item whose display name already resolves in the workspace is
      SKIPPED and its existing dataset id returned — re-runs never mint duplicate datasets
      (same philosophy as one-import-then-edit-in-place for workflows). The existing datasets come
      from `existing_json`, the MCP `search` result for `types: ["source"]`; without it the
      idempotency check cannot run and every item is treated as new, so re-runs CAN duplicate.
    - **Continue on error**: one bad file does not strand the rest; each item reports its own
      status, and the overall report separates created / skipped / failed.
    """
    if not isinstance(manifest, list) or not manifest:
        raise SavantAppApiError("Bulk dataset manifest must be a non-empty JSON list.")
    existing = {
        str(item.get("name", "")).strip().casefold(): item
        for item in load_records(existing_json, flag="--existing-json")
        if item.get("name")
    }
    created, skipped, failed = [], [], []
    for i, item in enumerate(manifest, 1):
        name = str(item.get("name") or "").strip()
        file_path = item.get("file")
        if not name or not file_path:
            failed.append({"index": i, "item": item, "error": "manifest item needs `file` and `name`"})
            continue
        prior = existing.get(name.casefold())
        if prior is not None:
            skipped.append({"index": i, "name": name, "datasetId": prior.get("id"),
                            "reason": "dataset with this name already exists in the workspace"})
            continue
        try:
            source = create_dataset(
                context, Path(file_path), name,
                item.get("type", "auto"),
                delimiter=item.get("delimiter", ","),
                sheet=item.get("sheet"),
                tab_pattern=item.get("tab_pattern"),
                sample_rows=int(item.get("sample_rows", 100)),
                skip_rows=int(item.get("skip_rows", 0)),
            )
            created.append({"index": i, "name": name, "datasetId": source.get("id"),
                            "status": source.get("status"), "source": source})
        except Exception as exc:  # noqa: BLE001 — per-item isolation is the contract
            failed.append({"index": i, "name": name, "file": str(file_path), "error": str(exc)})
    return {
        "task": "dataset_create_bulk",
        "requested": len(manifest),
        "created": created,
        "skipped": skipped,
        "failed": failed,
        "ok": not failed,
        # ready to paste into planner source plans: name -> dataset id for every resolved item
        "datasetIdsByName": {e["name"]: e["datasetId"] for e in created + skipped if e.get("datasetId")},
    }


def _safe_filename(value: str) -> str:
    cleaned = "".join("-" if char in '\\/:*?"<>|' else char for char in value).strip()
    return cleaned or "dataset"


def default_output_path(name: str) -> Path:
    return workspace_tmp("created-datasets", f"{_safe_filename(name)}.created-source.json")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Create Savant dataset(s) from local CSV / Excel / PDF files. Datasets are workspace-scoped and created in the authenticated session's namespace.")
    parser.add_argument("--file", help="Local file path (.csv/.txt, .xlsx/.xls, .pdf). Single-dataset mode.")
    parser.add_argument("--name", help="Dataset display name. Single-dataset mode.")
    parser.add_argument("--manifest",
                        help="Bulk mode: JSON file with a list of {file, name, type?, sheet?, tab_pattern?, "
                             "skip_rows?, delimiter?, sample_rows?} items. Creates all in one call; items whose "
                             "name already exists in the workspace are skipped (their existing id is returned); "
                             "one item's failure does not stop the rest.")
    parser.add_argument(
        "--type",
        default="auto",
        choices=["auto", "csv", "excel", "pdf"],
        help="File type; 'auto' detects from the extension.",
    )
    parser.add_argument("--sheet", help="Excel single-tab: sheet name (defaults to the first sheet).")
    parser.add_argument("--tab-pattern", help="Excel multi-tab: glob over sheet names to combine.")
    parser.add_argument("--skip-rows", type=int, default=0, help="Excel: number of leading rows to skip before the header row.")
    parser.add_argument("--delimiter", default=",")
    parser.add_argument("--sample-rows", type=int, default=100)
    parser.add_argument("--existing-json", type=Path,
                        help="MCP `search` result for types: [\"source\"]. Used by --manifest to skip "
                             "datasets whose display name already exists. Without it a re-run can "
                             "create duplicates.")
    parser.add_argument("--output-path", help="Optional path to write the created-source JSON.")
    args = parser.parse_args(argv)
    if bool(args.manifest) == bool(args.file or args.name):
        parser.error("Provide either --manifest (bulk) or --file/--name (single), not both.")
    if args.file and not args.name or args.name and not args.file:
        parser.error("Single-dataset mode requires both --file and --name.")

    context = api.discover_session(None)

    if args.manifest:
        manifest = json.loads(Path(args.manifest).read_text(encoding="utf-8"))
        report = create_datasets_bulk(context, manifest, existing_json=args.existing_json)
        for e in report["created"]:
            print(f"  [ok  ] created '{e['name']}' id={e['datasetId']} status={e.get('status')}")
        for e in report["skipped"]:
            print(f"  [skip] '{e['name']}' already exists, id={e['datasetId']}")
        for e in report["failed"]:
            print(f"  [FAIL] item {e['index']} '{e.get('name', '?')}': {e['error']}")
        out = Path(args.output_path) if args.output_path else default_output_path("bulk-manifest")
        save_json(report, out)
        print(f"BULK CREATE: {len(report['created'])} created, {len(report['skipped'])} skipped, "
              f"{len(report['failed'])} failed -> {out}")
        return 0 if report["ok"] else 1

    source = create_dataset(
        context,
        Path(args.file),
        args.name,
        args.type,
        delimiter=args.delimiter,
        sheet=args.sheet,
        tab_pattern=args.tab_pattern,
        sample_rows=args.sample_rows,
        skip_rows=args.skip_rows,
    )
    print(f"created dataset '{source.get('name')}' id={source.get('id')} status={source.get('status')}")
    save_json(source, Path(args.output_path) if args.output_path else default_output_path(args.name))
    return 0
